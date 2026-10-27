"""Formatting helpers of the handbook build (scripts/build_handbook.py): tables, the stage picture, the key-value and
verdict table looks."""

import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

def safe(name, limit=110):
    name = re.sub(r"\s*:\s+", " - ", name.replace("`", ""))
    name = re.sub(r"\s+/\s+", " - ", name)
    name = re.sub(r'[\\/:*?"<>|]', "-", name).strip().rstrip(".")
    return name[:limit].rstrip()

STAGE_ORDER = ["Pre-arrival", "Arrival", "Post-arrival", "Download", "Upload", "Pre-departure", "Departure"]

def stage_bar(stages, path, labels=None):
    """The seven turnaround stages in a row; the ones the module watches are filled."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle
    fig = plt.figure(figsize=(8.0, 0.5), dpi=220)
    ax = fig.add_axes([0, 0, 1, 1])
    n = len(STAGE_ORDER)
    for i, s in enumerate(STAGE_ORDER):
        on = s in stages
        ax.add_patch(Rectangle((i / n + 0.003, 0.08), 1 / n - 0.006, 0.84, facecolor="#1F5FBF" if on else "#E9EDF3",
                               edgecolor="none"))
        ax.text(i / n + 0.5 / n, 0.5, (labels or {}).get(s, s), ha="center", va="center", fontsize=8.5, family="DejaVu Sans",
                color="white" if on else "#5B6573", fontweight="bold" if on else "normal")
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    fig.savefig(path, facecolor="white")
    plt.close(fig)

def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:shd")):
        tcpr.remove(old)
    s = OxmlElement("w:shd")
    for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", fill)):
        s.set(qn(k), v)
    tcpr.append(s)

VERDICT_FILL = [("pass", "E2EFDA"), ("fail", "FADBD8"), ("not observed", "EDEDED")]

def polish(path, stage_png=None):
    """Reader-friendly touches on a converted page: the stage bar under the summary, a key-value look for the
    'At a glance' table (no empty header row, bold tinted labels) and coloured verdicts in the rule tables."""
    from docx.shared import Cm
    d = Document(path)
    if stage_png:
        anchor = next((p for p in d.paragraphs if p.style.name == "Block Text"), None) or d.paragraphs[min(1, len(d.paragraphs) - 1)]
        cap = d.add_paragraph()
        r = cap.add_run("Where in the turnaround it looks")
        r.font.size, r.font.color.rgb, r.italic = Pt(9), RGBColor(0x5B, 0x65, 0x73), True
        pic = d.add_paragraph()
        pic.add_run().add_picture(stage_png, width=Cm(16.5))
        anchor._p.addnext(cap._p)
        cap._p.addnext(pic._p)
    for t in d.tables:
        head = [c.text.strip().lower() for c in t.rows[0].cells]
        if len(head) == 2 and all(h == "" for h in head):
            t._tbl.remove(t.rows[0]._tr)
            for row in t.rows:
                shade(row.cells[0], "F2F5F9")
                for para in row.cells[0].paragraphs:
                    for run in para.runs:
                        run.bold = True
                row.cells[0].width, row.cells[1].width = Cm(4.2), Cm(12.3)
        elif head and head[0] == "verdict":
            for row in t.rows[1:]:
                v = row.cells[0].text.strip().lower()
                fill = "FFF2CC" if ("blocked" in v or "obstacle" in v) else next((f for k, f in VERDICT_FILL if v.startswith(k)), None)
                if fill:
                    shade(row.cells[0], fill)
                    for para in row.cells[0].paragraphs:
                        for run in para.runs:
                            run.bold = True
                row.cells[0].width, row.cells[1].width = Cm(4.2), Cm(12.3)
    d.save(path)

def style_tables(path):
    d = Document(path)
    if not d.tables:
        return
    for t in d.tables:
        pr = t._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", "A6A6A6")):
                e.set(qn(k), v)
            b.append(e)
        pr.append(b)
        for cell in t.rows[0].cells:
            shd = OxmlElement("w:shd")
            for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", "DCE6F2")):
                shd.set(qn(k), v)
            cell._tc.get_or_add_tcPr().append(shd)
    d.save(path)

def clean(md):
    md = md.lstrip()
    if md.startswith("# "):
        md = md.split("\n", 1)[1] if "\n" in md else ""
    md = re.sub(r"\[([^\]]+)\]\((?!https?://|mailto:)[^)]*\)", r"\1", md)  # repo-relative links -> text
    md = md.replace("@docs/", "docs/")
    return md.strip()

def md_table(header, rows):
    esc = lambda x: str(x).replace("|", "/").replace("\n", " ")
    lines = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    lines += ["| " + " | ".join(esc(x) for x in r) + " |" for r in rows]
    return "\n".join(lines)

def glance_of(md):
    sec = re.search(r"^## At a glance\s*\n(.*?)(?=^## )", md, flags=re.S | re.M)
    rows = {}
    for line in (sec.group(1).splitlines() if sec else []):
        c = [x.strip() for x in line.strip().strip("|").split("|")]
        if len(c) >= 2 and c[0] and not set(c[0]) <= set("-: "):
            rows[c[0]] = c[1]
    return rows
