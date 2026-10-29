"""Build the shared documentation folder (Word and Excel files) from the handbook in docs/handbook.

    python scripts/build_handbook.py --out G:/RampVision_Docs_v2

Every page of docs/handbook is a markdown file; this turns them into Word documents under the folder names people see:

    Start here.md                        -> Start here.docx
    teams/<team>/...                     -> 1 Teams/<Team>/...            (teams/cv/how-to -> 1 Teams/CV/How-to)
    system/...                           -> 2 System/...                  (system/shared -> 2 System/Shared parts)
    research/<team>/...                  -> 3 Research/<Team>/...

and adds what is computed: the owner and review line under every title (research reports carry their own date line), the
turnaround-stage picture and the latest monthly-test numbers on every check page, the numbers of the System overview,
the client rules (from the client's workbook through docs/05_module_logic.md) and the Modules index workbook. The check
pages follow the client's table: 31 checks under the client's names, in the table's seven stages
(docs/handbook/_data/modules.json, from `python scripts/handbook_module_data.py`; run it first after a new monthly
report).

Files edited by hand are never overwritten or removed: out/handbook/manifest.json keeps a checksum of every file this
build wrote, and a file that differs from it, a file open in Word or Excel, and every existing team page under 1 Teams
(the teams write their pages in Word) are kept as they are and listed at the end. Files of the previous build that are
no longer produced and were not edited are removed.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "scripts"))

import openpyxl  # noqa: E402
import pypandoc  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

from handbook_lib import STAGE_ORDER, clean, md_table, polish, safe, stage_bar, style_tables  # noqa: E402

HB = os.path.join(ROOT, "docs", "handbook")
REF = os.path.join(HB, "_build", "reference.docx")
WORK = os.path.join(ROOT, "out", "handbook")
TEAM_NAMES = {"cv": "CV", "mlops": "MLOps", "backend": "Backend", "data": "Data", "edge": "Edge"}
FOLDER_NAMES = {"architecture": "Architecture", "requirements": "Requirements", "shared": "Shared parts",
                "modules": "Modules", "how-to": "How-to"}
STAGE_LABELS = {"Pre-arrival": "Before arrival", "Arrival": "Arrival", "Post-arrival": "After the stop",
                "Download": "Unloading", "Upload": "Loading", "Pre-departure": "Before departure",
                "Departure": "Departure"}
OWNERS = [  # first match on the source path
    ("system/modules/", "CV team (owner of the check to assign)"),
    ("system/shared/Pose models", "Valentyn Fedorov (CV)"),
    ("system/shared/", "Yurii Luchko (CV)"),
    ("system/architecture/Edge", "Vladyslav Ivashchuk (Edge) with Maksym Chernyshev (MLOps)"),
    ("system/architecture/", "Maksym Chernyshev (MLOps)"),
    ("teams/cv/", "Valentyn Fedorov (CV team lead)"),
    ("teams/mlops/", "Maksym Chernyshev (MLOps)"),
    ("teams/backend/", "Aryan Singh (Backend)"),
    ("teams/data/", "Denys Filimonov (Data)"),
    ("teams/edge/", "Vladyslav Ivashchuk (Edge)"),
    ("", "Valentyn Fedorov (lead)"),
]


def target(rel: str) -> str:
    """docs/handbook-relative markdown path -> output-relative .docx path."""
    parts = rel[:-3].split("/")
    if parts[0] == "teams":
        parts = ["1 Teams", TEAM_NAMES.get(parts[1], parts[1])] + [FOLDER_NAMES.get(p, p) for p in parts[2:]]
    elif parts[0] == "system":
        parts = ["2 System"] + [FOLDER_NAMES.get(p, p) for p in parts[1:]]
    elif parts[0] == "research":
        parts = ["3 Research", TEAM_NAMES.get(parts[1], parts[1])] + parts[2:]
    return "/".join(parts) + ".docx"


def reviewed(rel: str) -> str:
    """The date of the last commit of the page; today for a page not committed yet or changed since."""
    path = os.path.join(HB, *rel.split("/"))
    dirty = subprocess.run(["git", "-C", ROOT, "status", "--porcelain", "--", path], capture_output=True, text=True).stdout
    last = subprocess.run(["git", "-C", ROOT, "log", "-1", "--format=%ad", "--date=short", "--", path],
                          capture_output=True, text=True).stdout.strip()
    return dt.date.today().isoformat() if dirty or not last else last


def sha256(path: str) -> str:
    with io.open(path, "rb") as fh:
        return hashlib.sha256(fh.read()).hexdigest()


class Build:
    def __init__(self, out: str, manifest: str):
        self.out = out
        self.manifest = manifest
        self.previous: dict = {}  # rel -> checksum of the file the last build wrote (None: an old manifest without them)
        self.previous_time = 0.0
        if os.path.exists(manifest):
            old = json.load(io.open(manifest, encoding="utf-8"))
            self.previous = old["files"] if isinstance(old, dict) else dict.fromkeys(old)
            self.previous_time = os.path.getmtime(manifest)
        self.made: list[str] = []
        self.locked: list[str] = []  # open in another program (Word, Excel): kept as they were
        self.kept: list[str] = []  # edited by hand since the last build: kept as they are
        self.teams: list[str] = []  # existing team pages: the teams keep them in Word

    def rel(self, path: str) -> str:
        return os.path.relpath(path, self.out).replace("\\", "/")

    def edited_by_hand(self, path: str) -> bool:
        rel = self.rel(path)
        if rel not in self.previous:
            return True  # not written by this build: someone else's file
        if self.previous[rel] is None:
            return os.path.getmtime(path) > self.previous_time + 60
        return sha256(path) != self.previous[rel]

    def keep_as_is(self, path: str) -> bool:
        """True (and the file is listed) when the build must not write it."""
        if not os.path.exists(path):
            return False
        try:
            io.open(path, "ab").close()
        except PermissionError:
            self.locked.append(path)
        else:
            if self.rel(path).startswith("1 Teams/"):
                self.teams.append(path)
            elif self.edited_by_hand(path):
                self.kept.append(path)
            else:
                return False
        self.made.append(path)
        return True

    def docx(self, rel_out: str, md: str, resdir: str | None = None) -> str | None:
        path = os.path.join(self.out, *rel_out.split("/"))
        os.makedirs(os.path.dirname(path), exist_ok=True)
        if self.keep_as_is(path):
            return None
        args = [f"--reference-doc={REF}"] + ([f"--resource-path={resdir}"] if resdir else [])
        pypandoc.convert_text(md.replace("<br>", "; "), "docx", format="gfm-raw_html", outputfile=path, extra_args=args)
        style_tables(path)
        self.made.append(path)
        return path

    def sheet(self, rel_out: str, title: str, header: list, rows: list, widths: list, links: dict) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = title
        ws.append(header)
        for r in rows:
            ws.append(list(r))
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E79")
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Arial")
                c.alignment = Alignment(vertical="top", wrap_text=True)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        for (r, c), link in links.items():
            cell = ws.cell(row=r, column=c)
            cell.hyperlink = link
            cell.font = Font(name="Arial", color="1F5FBF", underline="single")
        path = os.path.join(self.out, *rel_out.split("/"))
        os.makedirs(os.path.dirname(path), exist_ok=True)
        if self.keep_as_is(path):
            return
        wb.save(path)
        self.made.append(path)

    def copy(self, src: str, rel_out: str) -> None:
        path = os.path.join(self.out, *rel_out.split("/"))
        os.makedirs(os.path.dirname(path), exist_ok=True)
        if self.keep_as_is(path):
            return
        shutil.copy2(src, path)
        self.made.append(path)

    def finish(self) -> None:
        """Remove what the previous build made, this one did not and nobody edited; write the manifest."""
        now = {self.rel(p) for p in self.made}
        for rel in sorted(set(self.previous) - now):
            p = os.path.join(self.out, *rel.split("/"))
            if not os.path.exists(p):
                continue
            if rel.startswith("1 Teams/") or self.edited_by_hand(p):
                self.kept.append(p)
                now.add(rel)
                continue
            try:
                os.remove(p)
            except PermissionError:
                self.locked.append(p)
                now.add(rel)
        for dirpath, dirs, files in sorted(os.walk(self.out), reverse=True):  # empty folders left behind
            if dirpath != self.out and not dirs and not files:
                os.rmdir(dirpath)
        written = {}
        for rel in sorted(now):
            p = os.path.join(self.out, *rel.split("/"))
            if p in self.locked or p in self.kept or p in self.teams:  # not written now: remember whether edited
                try:
                    written[rel] = "edited by hand" if self.edited_by_hand(p) else self.previous.get(rel)
                except OSError:
                    written[rel] = "edited by hand"
            else:
                written[rel] = sha256(p)
        json.dump({"files": written}, io.open(self.manifest, "w", encoding="utf-8"), indent=1)


def with_meta(md: str, rel: str) -> str:
    """The owner and review line under the title; research reports have their own date line."""
    if rel.startswith("research/"):
        return md
    owner = next(o for prefix, o in OWNERS if rel.startswith(prefix))
    meta = f"*Owner: {owner} · Last reviewed: {reviewed(rel)} · Next review: first week of next month*"
    m = re.match(r"\s*(# .+\n)", md)
    return (m.group(1) + "\n" + meta + "\n" + md[m.end():]) if m else meta + "\n\n" + md


def accuracy_rows(check: dict) -> tuple[str, str]:
    t = check["monthly_test"]
    if not t:
        return "no verdict on the monthly test", "—"
    acc = f"{t['correct']} of {t['labelled']} turnarounds correct ({t['percent']} %) on the monthly test"
    fails = (f"{t['fails_caught']} of {t['fails_labelled']}" if t["fails_labelled"]
             else "no turnaround in the monthly test was labelled Fail")
    return acc, fails


def module_cell(check: dict) -> str:
    if check["derived_from"]:
        return "none: combined from " + " and ".join(check["derived_from"])
    if not check["production_version"]:
        return f"{check['repository']} (does not run in production)"
    others = check["other_checks_of_module"]
    return check["repository"] + (f" (also gives: {'; '.join(others)})" if others else "")


def fresh_status(md: str, check: dict) -> str:
    """Replace the two accuracy rows of the Status block with the latest monthly-test numbers."""
    acc, fails = accuracy_rows(check)
    md = re.sub(r"^\| How accurate it is \|.*$", lambda _: f"| How accurate it is | {acc} |", md, count=1, flags=re.M)
    return re.sub(r"^\| Fail turnarounds it caught \|.*$", lambda _: f"| Fail turnarounds it caught | {fails} |", md,
                  count=1, flags=re.M)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", required=True, help="the shared folder to write, e.g. G:/RampVision_Docs_v2")
    a = ap.parse_args()
    os.makedirs(WORK, exist_ok=True)
    b = Build(os.path.abspath(a.out), os.path.join(WORK, "manifest.json"))
    data = json.load(io.open(os.path.join(HB, "_data", "modules.json"), encoding="utf-8"))
    by_output = {c["output"]: c for c in data["checks"]}

    # every handbook page
    for dirpath, _dirs, files in os.walk(HB):
        rel_dir = os.path.relpath(dirpath, HB).replace("\\", "/")
        if rel_dir.split("/")[0] in ("_data", "_build"):
            continue
        for f in sorted(files):
            if not f.endswith(".md") or f == "STYLE.md":
                continue
            rel = f if rel_dir == "." else f"{rel_dir}/{f}"
            md = io.open(os.path.join(dirpath, f), encoding="utf-8").read()
            check = by_output.get(f"docs/handbook/{rel}")
            if check:
                md = fresh_status(md, check)
            if rel == "system/System overview.md":
                md = overview_numbers(md, data)
            out = b.docx(target(rel), with_meta(md, rel), resdir=dirpath)
            if out is None:
                continue
            if check:
                png = os.path.join(WORK, safe(check["file_name"]) + "_stages.png")
                stage_bar(check["stages"], png, STAGE_LABELS)
                polish(out, png)
            elif rel.startswith("system/"):
                polish(out)

    # client rules: generated from the client's workbook
    rules = clean(io.open(os.path.join(ROOT, "docs", "05_module_logic.md"), encoding="utf-8").read())
    b.docx("2 System/Requirements/Client rules.docx",
           "# Client rules — Pass, Fail and Not observed for every check\n\n"
           "*Owner: Valentyn Fedorov (lead) with Oksana (QA) · Source: the client's workbook next to this page*\n\n" + rules)
    b.copy(os.path.join(ROOT, "docs", "DX_Modules_logic.xlsx"), "2 System/Requirements/Client rules (client original).xlsx")

    # the index of the checks, in the order of the client's table
    rows, links = [], {}
    for c in sorted(data["checks"], key=lambda c: c["order"]):
        acc, fails = accuracy_rows(c)
        page = target(c["output"][len("docs/handbook/"):])
        real_time = "follows its two parts" if c["derived_from"] else re.sub(r"[`*]", "", c["glance"].get("Real time", ""))
        rows.append([c["stage"], c["name"], c["plain_name"], module_cell(c), ", ".join(c["stages_plain"]), acc, fails,
                     real_time, "to assign", "open"])
        links[(len(rows) + 1, 10)] = os.path.relpath(os.path.join(b.out, *page.split("/")),
                                                     os.path.join(b.out, "2 System", "Modules")).replace("\\", "/")
    b.sheet("2 System/Modules/Modules index.xlsx", "Checks",
            ["Stage", "Check (the client's name)", "In plain words", "Module", "Stages it looks at",
             "Accuracy on the monthly test", "Fail turnarounds caught", "Real time", "Who looks after it", "Page"],
            rows, [22, 46, 38, 40, 30, 40, 22, 40, 14, 8], links)

    b.finish()
    print(f"{len(b.made)} files in {b.out}")
    for p in b.locked:
        print(f"not updated, open in another program: {b.rel(p)}")
    for p in b.kept:
        print(f"kept as edited by hand: {b.rel(p)}")
    if b.teams:
        print(f"{len(b.teams)} existing team pages kept as they are (the teams edit them in Word)")
    return 0


def overview_numbers(md: str, data: dict) -> str:
    scored = sorted(((c["monthly_test"]["percent"], c["name"]) for c in data["checks"] if c["monthly_test"]),
                    key=lambda x: x[0])
    lo, hi = scored[0][0], scored[-1][0]
    weakest = ", ".join(f"*{name}* ({p} %)" for p, name in scored[:5])
    text = (f"On the monthly test ({data['monthly_test']['turnarounds']} turnarounds the client checked by hand) the checks "
            f"give the right verdict in **{lo} to {hi} %** of the turnarounds. The weakest: {weakest}. Per check, with the "
            "turnarounds labelled Fail that each one caught: *Modules / Modules index*.")
    return md.replace("{{ACCURACY_RANGE}}", f"{lo}–{hi} %").replace("{{ACCURACY_TEXT}}", text)


if __name__ == "__main__":
    raise SystemExit(main())
